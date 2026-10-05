#!/usr/bin/env python3
"""
Seguimiento Post-CAPA · Blue Medical
Convierte el Excel de formadores en el archivo cifrado `datos.enc.json` que lee la página.

Uso:
    python3 actualizar_datos.py Seguimiento_Post-CAPA_formadores.xlsx [salida.json]
La contraseña se pide por consola (o por la variable de entorno POSTCAPA_CLAVE).
Requiere: openpyxl, cryptography.
"""
import sys, os, json, re, base64, secrets, getpass, datetime as dt, unicodedata
from openpyxl import load_workbook
from cryptography.hazmat.primitives.ciphers.aead import AESGCM
from cryptography.hazmat.primitives.kdf.pbkdf2 import PBKDF2HMAC
from cryptography.hazmat.primitives import hashes

ITER = 250_000
IGNORAR = {"instrucciones", "plantilla", "listas"}
CRITERIOS = [("dom", "Dominio de procesos"), ("sis", "Manejo de sistemas"), ("pac", "Atención al paciente"),
             ("ven", "Venta / orientación"), ("seg", "Seguridad al ejecutar")]
CHECK = {"✔", "✓", "x", "X", "si", "sí", "Si", "Sí", "SI", "SÍ", True, 1, "TRUE", "VERDADERO"}


def txt(v):
    if v is None: return ""
    if isinstance(v, float) and v.is_integer(): v = int(v)
    return str(v).strip()


def fecha(v):
    if v in (None, ""): return ""
    if isinstance(v, dt.datetime): return v.date().isoformat()
    if isinstance(v, dt.date): return v.isoformat()
    m = re.match(r"^\s*(\d{1,2})[/-](\d{1,2})[/-](\d{2,4})\s*$", str(v))
    if m:
        d, mth, y = map(int, m.groups()); y = y + 2000 if y < 100 else y
        return dt.date(y, mth, d).isoformat()
    return ""


def entero(v):
    if v in (None, ""): return None
    try: return max(0, int(round(float(v))))
    except (TypeError, ValueError): return None


def slug(s):
    s = unicodedata.normalize("NFD", s).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]+", "-", s).strip("-")


def find_row(ws, text, col=1, start=1, end=200):
    for r in range(start, end):
        if txt(ws.cell(r, col).value).lower().startswith(text.lower()): return r
    return None


def leer(path):
    wb = load_workbook(path, data_only=True)
    temas = [txt(c.value) for c in wb["Listas"]["D"][1:] if txt(c.value)] if "Listas" in wb.sheetnames else []
    avisos, grupos, colabs = [], [], []
    for ws in wb.worksheets:
        if ws.title.strip().lower() in IGNORAR: continue
        if txt(ws["A3"].value).lower() != "fecha de ingreso":
            avisos.append(f"Hoja «{ws.title}» omitida: no tiene el formato de la plantilla."); continue
        ingreso, fin, formador = fecha(ws["B3"].value), fecha(ws["D3"].value), txt(ws["F3"].value)
        if not fin:
            avisos.append(f"Hoja «{ws.title}»: falta la fecha de fin de CAPA; se omitió."); continue
        grupos.append({"hoja": ws.title, "ingreso": ingreso, "finCapa": fin, "formador": formador})
        # encabezados de la sección 2
        hr = find_row(ws, "Colaborador", 1, 15)
        heads = {txt(ws.cell(hr, j).value).lower(): j for j in range(1, ws.max_column + 1)}
        def col(name):
            for k, j in heads.items():
                if k.startswith(name.lower()): return j
            return None
        c_obs = col("¿cómo le fue"); c_otro = col("otro tema"); c_det = col("detalle"); c_acc = col("¿qué se acordó")
        c_resp = col("responsable"); c_aseg = col("asegurado"); c_noas = col("no asegurado"); c_adic = col("cantidad de adicionales")
        c_com = col("comentarios"); c_freal = col("fecha realizada")
        t0 = (c_obs or 12) + 1
        tema_cols = [(j, temas[j - t0] if j - t0 < len(temas) else txt(ws.cell(hr, j).value)) for j in range(t0, c_otro or t0)]
        # colaboradores (sección 1)
        h1 = find_row(ws, "Nombre completo", 1, 5, hr)
        personas, orden = {}, []
        for r in range(h1 + 1, hr - 4):          # cupos de la sección 1 (en orden)
            nombre = txt(ws.cell(r, 1).value)
            orden.append(nombre)
            if nombre:
                falt = [n for n, j in (("puesto", 2), ("sede", 3), ("jefe", 4)) if not txt(ws.cell(r, j).value)]
                if falt: avisos.append(f"«{nombre}» ({ws.title}): falta {', '.join(falt)}.")
                personas[nombre] = {"id": slug(nombre + "-" + fin), "nombre": nombre, "puesto": txt(ws.cell(r, 2).value),
                                    "sede": txt(ws.cell(r, 3).value), "jefe": txt(ws.cell(r, 4).value), "correo": txt(ws.cell(r, 5).value),
                                    "ingreso": ingreso, "finCapa": fin, "formador": formador, "grupo": ws.title,
                                    "semanas": [None] * 4}
        # filas semanales: 4 filas por cupo, en el mismo orden que la sección 1
        s0 = hr + 1
        for i, nombre in enumerate(orden):
            if not nombre: continue
            for k in range(4):
                r = s0 + i * 4 + k; n = k + 1
                scores = {}
                for q, (cid, cn) in enumerate(CRITERIOS):
                    v = entero(ws.cell(r, 5 + q).value)
                    if v is not None:
                        if 1 <= v <= 5: scores[cid] = v
                        else: avisos.append(f"«{nombre}» semana {n}: calificación fuera de rango en {cn} ({v}).")
                temas_m = [t for j, t in tema_cols if t and (ws.cell(r, j).value in CHECK or txt(ws.cell(r, j).value) in CHECK)]
                g = lambda c: ws.cell(r, c).value if c else None
                sem = {"n": n, "fechaReal": fecha(g(c_freal)), "scores": scores, "obs": txt(g(c_obs)), "temas": temas_m,
                       "otroTema": txt(g(c_otro)), "detalle": txt(g(c_det)), "accion": txt(g(c_acc)), "responsable": txt(g(c_resp)),
                       "aseg": entero(g(c_aseg)), "noaseg": entero(g(c_noas)), "adic": entero(g(c_adic)), "comentarios": txt(g(c_com))}
                if sem["fechaReal"] == "" and scores:
                    avisos.append(f"«{nombre}» semana {n}: tiene calificaciones pero no fecha realizada; se usará la fecha programada.")
                personas[nombre]["semanas"][k] = sem
        for p in personas.values():
            p["semanas"] = [s or {"n": i + 1, "fechaReal": "", "scores": {}, "obs": "", "temas": [], "otroTema": "", "detalle": "", "accion": "",
                                  "responsable": "", "aseg": None, "noaseg": None, "adic": None, "comentarios": ""} for i, s in enumerate(p["semanas"])]
            colabs.append(p)
    data = {"version": 2, "generado": dt.datetime.now().isoformat(timespec="minutes"),
            "criterios": [{"id": a, "nombre": b} for a, b in CRITERIOS], "temas": temas, "grupos": grupos, "colaboradores": colabs}
    return data, avisos


def cifrar(data, clave):
    salt, iv = secrets.token_bytes(16), secrets.token_bytes(12)
    key = PBKDF2HMAC(algorithm=hashes.SHA256(), length=32, salt=salt, iterations=ITER).derive(clave.encode("utf-8"))
    ct = AESGCM(key).encrypt(iv, json.dumps(data, ensure_ascii=False).encode("utf-8"), None)
    b = lambda x: base64.b64encode(x).decode()
    return {"formato": "postcapa-aesgcm-v1", "iter": ITER, "salt": b(salt), "iv": b(iv), "datos": b(ct)}


if __name__ == "__main__":
    if len(sys.argv) < 2: sys.exit(__doc__)
    data, avisos = leer(sys.argv[1])
    out = sys.argv[2] if len(sys.argv) > 2 else "datos.enc.json"
    clave = os.environ.get("POSTCAPA_CLAVE") or getpass.getpass("Contraseña del dashboard: ")
    json.dump(cifrar(data, clave), open(out, "w"), ensure_ascii=False)
    print(f"{len(data['colaboradores'])} colaboradores en {len(data['grupos'])} grupo(s) → {out}")
    for a in avisos: print("AVISO:", a)
